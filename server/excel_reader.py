from __future__ import annotations

from io import BytesIO
from tempfile import NamedTemporaryFile
from typing import Any

import xlrd
from fastapi import HTTPException
from openpyxl import load_workbook
from pyxlsb import open_workbook as open_xlsb_workbook

SUPPORTED_EXCEL_EXTENSIONS = (
    ".xlsx",
    ".xls",
    ".xlsm",
    ".xlsb",
    ".xltx",
    ".xltm",
    ".xlam",
)

_OPENPYXL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xlam"}
_XLRD_EXTENSIONS = {".xls"}
_PYXLSB_EXTENSIONS = {".xlsb"}


def _extract_extension(filename: str | None) -> str:
    name = (filename or "").strip().lower()
    if "." not in name:
        return ""
    return f".{name.rsplit('.', 1)[-1]}"


def ensure_supported_excel_filename(filename: str | None):
    extension = _extract_extension(filename)
    if extension not in SUPPORTED_EXCEL_EXTENSIONS:
        expected = ", ".join(SUPPORTED_EXCEL_EXTENSIONS)
        raise HTTPException(status_code=400, detail=f"Unsupported file type, expected one of: {expected}")


def parse_excel_bytes(raw: bytes, filename: str | None) -> dict[str, list[list[Any]]]:
    extension = _extract_extension(filename)
    ensure_supported_excel_filename(filename)

    try:
        if extension in _OPENPYXL_EXTENSIONS:
            workbook = load_workbook(filename=BytesIO(raw), read_only=True, data_only=True)
            data: dict[str, list[list[Any]]] = {}
            for sheet in workbook.worksheets:
                data[sheet.title] = [list(row) for row in sheet.iter_rows(values_only=True)]
            return data

        if extension in _XLRD_EXTENSIONS:
            workbook = xlrd.open_workbook(file_contents=raw)
            data = {}
            for sheet in workbook.sheets():
                rows: list[list[Any]] = []
                for row_index in range(sheet.nrows):
                    rows.append([sheet.cell_value(row_index, col_index) for col_index in range(sheet.ncols)])
                data[sheet.name] = rows
            return data

        if extension in _PYXLSB_EXTENSIONS:
            data = {}
            with NamedTemporaryFile(delete=True, suffix=extension) as temp:
                temp.write(raw)
                temp.flush()
                with open_xlsb_workbook(temp.name) as workbook:
                    for sheet_name in workbook.sheets:
                        rows: list[list[Any]] = []
                        with workbook.get_sheet(sheet_name) as sheet:
                            for row in sheet.rows():
                                values = [cell.v for cell in row]
                                while values and values[-1] is None:
                                    values.pop()
                                rows.append(values)
                        data[sheet_name] = rows
            return data

    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"Failed to parse workbook: {error}") from error

    raise HTTPException(status_code=400, detail="Unsupported workbook format")
